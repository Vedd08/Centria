import openpyxl
from collections import Counter
wb = openpyxl.load_workbook(r'C:/Users/Vedant/Downloads/Data Entry.xlsx', data_only=True)

def vals(ws, col, start):
    return [str(ws.cell(r,col).value).strip() if ws.cell(r,col).value is not None else '' for r in range(start, ws.max_row+1)]

opd, ipd = wb['Daily Diary OPD'], wb['Daily Diary IPD']
pm = wb['Patient Master']

for label, ws, start in [('OPD', opd, 3), ('IPD', ipd, 2)]:
    dates = vals(ws,1,start); names = vals(ws,3,start); mob = vals(ws,2,start)
    tr = vals(ws,5,start); rm = vals(ws,11,start); op = vals(ws,12,start)
    cash = vals(ws,16,start); bank = vals(ws,17,start)
    n = sum(1 for d in dates if d)
    print(f'{label}: rows={n}')
    print(f'   blank mobile: {sum(1 for i,d in enumerate(dates) if d and not mob[i])}')
    print(f'   blank name  : {sum(1 for i,d in enumerate(dates) if d and not names[i])}')
    print(f'   blank treat : {sum(1 for i,d in enumerate(dates) if d and not tr[i])}')
    print(f'   blank RM    : {sum(1 for i,d in enumerate(dates) if d and not rm[i])}')
    print(f'   blank oper  : {sum(1 for i,d in enumerate(dates) if d and not op[i])}')
    print(f'   has cash    : {sum(1 for i,d in enumerate(dates) if d and cash[i] and cash[i]!="0.0")}')
    print(f'   has bank    : {sum(1 for i,d in enumerate(dates) if d and bank[i] and bank[i]!="0.0")}')
    trs = set(t for t in tr if t); print(f'   distinct treatments used: {len(trs)}')

# patient master
pnames = vals(pm,3,2); pmob = vals(pm,4,2)
print(f'\nPatient Master rows: {sum(1 for n in pnames if n)}  distinct mobiles: {len(set(m for m in pmob if m))}')
# do diary patients exist in master?
master_mob = set(m.replace(".0","") for m in pmob if m)
for label, ws, start in [('OPD', opd, 3), ('IPD', ipd, 2)]:
    mob = [m.replace(".0","") for m in vals(ws,2,start) if m]
    missing = set(mob) - master_mob
    print(f'{label}: diary mobiles not in Patient Master: {len(missing)} / {len(set(mob))}')
wb.close()
