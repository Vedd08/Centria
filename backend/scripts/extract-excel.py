"""
Flattens the clinic's "Data Entry.xlsx" workbook into a single JSON file that
prisma/import-excel.ts can load.

Kept as a separate step so the (slow, dependency-heavy) spreadsheet parsing is
decoupled from the database write, and so the intermediate JSON can be inspected
before anything touches the database.

Usage:
    pip install openpyxl
    python scripts/extract-excel.py "C:/Users/<you>/Downloads/Data Entry.xlsx" C:/tmp/hms-import.json
"""

import datetime
import json
import sys

import openpyxl

SRC = sys.argv[1] if len(sys.argv) > 1 else r"C:/Users/Vedant/Downloads/Data Entry.xlsx"
DEST = sys.argv[2] if len(sys.argv) > 2 else r"C:/tmp/hms-import.json"


def s(v):
    """Cell -> trimmed string. Dates become YYYY-MM-DD; whole floats lose the '.0'."""
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def num(v):
    try:
        if v is None or str(v).strip() == "":
            return 0.0
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def phone(v):
    """Mobile number is the clinic's real patient key, so normalise it hard."""
    p = s(v).replace(".0", "").replace(" ", "").replace("-", "").replace("+91", "")
    return p if p and p.lower() not in ("none", "nan", "#n/a") else ""


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    out = {}

    # --- Patient Master ---
    pm = wb["Patient Master"]
    patients = {}
    for r in range(2, pm.max_row + 1):
        name = s(pm.cell(r, 3).value)
        if not name:
            continue
        ph = phone(pm.cell(r, 4).value)
        patients[ph or f"nophone::{name.lower()}"] = {
            "name": name,
            "phone": ph,
            "email": s(pm.cell(r, 5).value),
            "dob": s(pm.cell(r, 6).value),
            "gender": s(pm.cell(r, 7).value) or "Other",
            "address": s(pm.cell(r, 8).value),
            "allergyNotes": s(pm.cell(r, 9).value),
            "howHeardAboutUs": s(pm.cell(r, 10).value),
            "referringDoctor": s(pm.cell(r, 11).value),
        }
    out["patients"] = patients

    # --- Treatment Master ---
    tm = wb["Treatment Master"]
    out["treatments"] = {
        s(tm.cell(r, 1).value): s(tm.cell(r, 2).value).lower() == "yes"
        for r in range(2, tm.max_row + 1)
        if s(tm.cell(r, 1).value)
    }

    # --- Day books. Headers sit on different rows in each sheet. ---
    visits = []
    for vtype, sheet, start in [("OPD", "Daily Diary OPD", 3), ("IPD", "Daily Diary IPD", 2)]:
        ws = wb[sheet]
        for r in range(start, ws.max_row + 1):
            d = s(ws.cell(r, 1).value)
            if not d or not d.startswith("20"):
                continue
            nm, ph = s(ws.cell(r, 3).value), phone(ws.cell(r, 2).value)
            if not nm and not ph:
                continue
            visits.append({
                "visitType": vtype,
                "date": d,
                "phone": ph,
                "patientName": nm,
                "diagnosis": s(ws.cell(r, 4).value),
                "treatment": s(ws.cell(r, 5).value),
                "remarks": s(ws.cell(r, 6).value),
                "status": s(ws.cell(r, 7).value) or "New",
                "taxable": s(ws.cell(r, 8).value).lower() == "yes",
                "nextAppointmentDate": s(ws.cell(r, 9).value),
                "doctorName": s(ws.cell(r, 10).value),
                "clientRm": s(ws.cell(r, 11).value),
                "operator": s(ws.cell(r, 12).value),
                "otherRemarks": s(ws.cell(r, 13).value),
                "cash": num(ws.cell(r, 16).value),
                "bank": num(ws.cell(r, 17).value),
                "paymentRefId": s(ws.cell(r, 18).value),
            })
    out["visits"] = visits

    # --- Package Details ---
    pd_ = wb["Package Details"]
    out["packages"] = [
        {
            "phone": phone(pd_.cell(r, 2).value),
            "patientName": s(pd_.cell(r, 3).value),
            "purchaseDate": s(pd_.cell(r, 4).value),
            "code": s(pd_.cell(r, 5).value),
            "amount": num(pd_.cell(r, 6).value),
            "cash": num(pd_.cell(r, 7).value),
            "bank": num(pd_.cell(r, 8).value),
            "paymentRefId": s(pd_.cell(r, 9).value),
            "description": s(pd_.cell(r, 10).value),
            "taxable": s(pd_.cell(r, 11).value).lower() == "yes",
            "totalSittings": int(num(pd_.cell(r, 12).value)),
            "availed": int(num(pd_.cell(r, 13).value)),
            "available": int(num(pd_.cell(r, 14).value)),
            "status": s(pd_.cell(r, 15).value) or "Follow Up",
        }
        for r in range(3, pd_.max_row + 1)
        if s(pd_.cell(r, 3).value)
    ]

    # --- Purchase Register ---
    pr = wb["Purchase Register"]
    out["purchases"] = [
        {
            "vendorName": s(pr.cell(r, 1).value),
            "vendorGstn": s(pr.cell(r, 2).value) or "UNKNOWN",
            "invoiceNumber": s(pr.cell(r, 3).value) or f"AUTO-{r}",
            "invoiceDate": s(pr.cell(r, 4).value),
            "inputClassification": s(pr.cell(r, 5).value) or "Other",
            "itcEligibility": s(pr.cell(r, 6).value) or "Ineligible",
            "taxableValue": num(pr.cell(r, 7).value),
            "igst": num(pr.cell(r, 8).value),
            "cgst": num(pr.cell(r, 9).value),
            "sgst": num(pr.cell(r, 10).value),
            "totalValue": num(pr.cell(r, 11).value),
        }
        for r in range(3, pr.max_row + 1)
        if s(pr.cell(r, 1).value)
    ]

    # --- Doctors / Operators master ---
    dm = wb["Doctors Masters"]
    out["doctors"] = sorted({s(dm.cell(r, 3).value) for r in range(3, dm.max_row + 1) if s(dm.cell(r, 3).value)})
    out["operators"] = sorted({s(dm.cell(r, 1).value) for r in range(3, dm.max_row + 1) if s(dm.cell(r, 1).value)})

    with open(DEST, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False)

    print(f"patients   : {len(patients)}")
    print(f"treatments : {len(out['treatments'])}")
    print(f"visits     : {len(visits)} (OPD {sum(1 for v in visits if v['visitType'] == 'OPD')}"
          f" / IPD {sum(1 for v in visits if v['visitType'] == 'IPD')})")
    print(f"packages   : {len(out['packages'])}")
    print(f"purchases  : {len(out['purchases'])}")
    print(f"-> {DEST}")
    wb.close()


if __name__ == "__main__":
    main()
