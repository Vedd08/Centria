import openpyxl
from collections import Counter
wb = openpyxl.load_workbook(r'C:/Users/Vedant/Downloads/Data Entry.xlsx', data_only=True)

def vals(ws,col,start):
    return [str(ws.cell(r,col).value).strip() if ws.cell(r,col).value is not None else '' for r in range(start, ws.max_row+1)]

# Other Income - any real data?
oi = wb['Other Income']
nonempty = 0
for r in range(2, oi.max_row+1):
    row = [oi.cell(r,c).value for c in range(2,8)]  # skip timestamp
    if any(v is not None and str(v).strip() and str(v).strip()!='0' for v in row):
        nonempty += 1
        if nonempty <= 5:
            print('OtherIncome sample:', [str(v)[:20] if v is not None else '' for v in [oi.cell(r,c).value for c in range(1,11)]])
print(f'Other Income rows with real data: {nonempty}')

# Purchase register values
pr = wb['Purchase Register']
print('\nPurchase ITC Availability:', dict(Counter(v for v in vals(pr,6,3) if v).most_common()))
print('Purchase Input Classification:', dict(Counter(v for v in vals(pr,5,3) if v).most_common()))
print('Purchase rows:', sum(1 for v in vals(pr,1,3) if v))

# Appointment Due
ad = wb['Appointment Due']
n=0
for r in range(6, ad.max_row+1):
    if any(ad.cell(r,c).value for c in range(1,6)):
        n+=1
        if n<=3: print('ApptDue sample:', [str(ad.cell(r,c).value)[:22] if ad.cell(r,c).value else '' for c in range(1,6)])
print(f'Appointment Due data rows: {n}')

# Package details full columns
pd_ = wb['Package Details']
print('\nPackage headers:', [pd_.cell(2,c).value for c in range(1,16)])
print('Package status values:', dict(Counter(v for v in vals(pd_,15,3) if v).most_common()))

# Patient master how-you-know + referring
pm = wb['Patient Master']
print('\nHow you know us:', dict(Counter(v for v in vals(pm,10,2) if v).most_common(8)))
print('Gender:', dict(Counter(v for v in vals(pm,7,2) if v).most_common()))
wb.close()
